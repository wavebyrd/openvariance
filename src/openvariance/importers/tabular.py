"""Read messy spreadsheet exports into clean rows, without guessing quietly.

Handles the layouts accounting systems actually produce:

* **wide**: accounts down the side, periods across the top (QuickBooks, Xero
  and NetSuite P&L exports, most budget spreadsheets)
* **long**: one row per account and period, optionally with department and
  debit/credit columns (GL detail and trial balance exports)
* **by department**: accounts down the side, departments or classes across
  the top, for a single period you name

Along the way it skips titles, blank rows and section headers, drops subtotal
and total rows (verifying each against the rows it summarizes), applies
"in thousands" style units, and parses accounting number formats such as
``(1,234.50)`` and ``1,234.50-``. Everything it does is recorded in the
:class:`Inspection` so a person or an agent can review it.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
from collections import defaultdict
from collections.abc import Sequence
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from ..money import ZERO

Cell = Any

MONTHS = {
    m: i
    for i, names in enumerate(
        [
            ("jan", "january"),
            ("feb", "february"),
            ("mar", "march"),
            ("apr", "april"),
            ("may",),
            ("jun", "june"),
            ("jul", "july"),
            ("aug", "august"),
            ("sep", "sept", "september"),
            ("oct", "october"),
            ("nov", "november"),
            ("dec", "december"),
        ],
        start=1,
    )
    for m in names
}
_MONTH_RE = "|".join(sorted(MONTHS, key=len, reverse=True))

TOTAL_WORDS = re.compile(
    r"^(total\b|net\s+(income|loss|profit|operating income|ordinary income|other income|"
    r"other expense|revenue|sales)|gross\s+(profit|margin)|operating\s+(income|profit|loss)|"
    r"ebitda\b|ebit\b|income\s+before|profit\s+before|earnings\s+before)",
    re.IGNORECASE,
)
UNIT_PATTERNS = (
    (
        re.compile(
            r"in\s+millions|\(\s*\$?\s*mm?\s*\)|\$\s*mm?\b|\(\s*\$?\s*000,?000s?\s*\)",
            re.IGNORECASE,
        ),
        Decimal(1_000_000),
    ),
    (
        re.compile(
            r"in\s+thousands|\(\s*\$?\s*000s?\s*\)|\$\s*000s?\b|\$k\b|\(\s*\$?\s*k\s*\)",
            re.IGNORECASE,
        ),
        Decimal(1_000),
    ),
)
ACCOUNT_NUMBER = re.compile(r"^\s*(\d{3,6}(?:[.\-]\d{1,4})?)\s*(?:[-–·:.|]\s*|\s+)(.+?)\s*$")


@dataclass
class SourceRow:
    row: int  # 1-based row number in the source sheet
    label: str
    number: str | None
    name: str
    values: dict[str, Decimal]  # column key (period or department) -> amount, units applied
    kind: str  # "detail", "total", "subtotal"
    section: str = ""
    department: str | None = None
    note: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "row": self.row,
            "label": self.label,
            "kind": self.kind,
            "section": self.section,
            "values": {k: str(v) for k, v in self.values.items()},
            **({"department": self.department} if self.department else {}),
            **({"note": self.note} if self.note else {}),
        }


@dataclass
class Inspection:
    path: str
    sheet: str | None
    layout: str  # "wide", "long", "by_department"
    header_row: int
    columns: list[str]  # periods for wide/long, departments for by_department
    units: Decimal = Decimal(1)
    rows: list[SourceRow] = field(default_factory=list)
    skipped: list[tuple[int, str]] = field(default_factory=list)
    hints: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def details(self) -> list[SourceRow]:
        return [r for r in self.rows if r.kind == "detail"]

    @property
    def totals(self) -> list[SourceRow]:
        return [r for r in self.rows if r.kind in ("total", "subtotal")]

    @property
    def periods(self) -> list[str]:
        return [] if self.layout == "by_department" else self.columns

    def to_dict(self, max_rows: int = 200) -> dict[str, Any]:
        return {
            "file": self.path,
            "sheet": self.sheet,
            "layout": self.layout,
            "header_row": self.header_row,
            "columns": self.columns,
            "units_multiplier": str(self.units),
            "detail_rows": len(self.details),
            "total_rows": len(self.totals),
            "rows": [r.to_dict() for r in self.rows[:max_rows]],
            "skipped": [{"row": n, "reason": why} for n, why in self.skipped[:max_rows]],
            "hints": self.hints,
            "warnings": self.warnings,
        }


# Cell parsing -------------------------------------------------------------


def parse_amount(cell: Cell) -> Decimal | None:
    """Parse an accounting-formatted amount. Returns None for blanks and dashes."""
    if cell is None or isinstance(cell, bool):
        return None
    if isinstance(cell, (int, Decimal)):
        return Decimal(cell)
    if isinstance(cell, float):
        return Decimal(repr(cell))
    if not isinstance(cell, str):
        return None
    text = cell.strip()
    if text in ("", "-", "–", "—", "n/a", "N/A", "nm", "n/m"):
        return None
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative, text = True, text[1:-1]
    if text.endswith("-") and len(text) > 1:
        negative, text = True, text[:-1]
    if text.upper().endswith(" CR"):
        negative, text = True, text[:-3]
    elif text.upper().endswith(" DR"):
        text = text[:-3]
    text = re.sub(r"^[A-Z]{3}\s+|[$€£¥\s,]", "", text)
    if text.startswith("-"):
        negative, text = not negative, text[1:]
    if not re.fullmatch(r"\d+(\.\d+)?|\.\d+", text):
        return None
    try:
        value = Decimal(text)
    except InvalidOperation:
        return None
    return -value if negative else value


def parse_period(cell: Cell) -> str | None:
    """Recognize a month or quarter header. Returns ``YYYY-MM`` or ``YYYY-Qn``."""
    if isinstance(cell, (dt.datetime, dt.date)):
        return f"{cell.year:04d}-{cell.month:02d}"
    if not isinstance(cell, str):
        return None
    text = cell.strip().lower().replace(".", "")
    if not text or re.search(r"\b(total|ytd|to date)\b", text) or re.search(r"\s[-–]\s", text):
        return None
    if m := re.fullmatch(r"(\d{4})[-/](\d{1,2})(?:[-/]\d{1,2})?", text):
        return _ym(int(m[1]), int(m[2]))
    if m := re.fullmatch(r"(\d{1,2})[-/](\d{4})", text):
        return _ym(int(m[2]), int(m[1]))
    if m := re.fullmatch(rf"({_MONTH_RE})[\s\-',]*(\d{{4}}|\d{{2}})", text):
        return _ym(_year(m[2]), MONTHS[m[1]])
    if m := re.fullmatch(rf"(\d{{4}})[\s\-]*({_MONTH_RE})", text):
        return _ym(int(m[1]), MONTHS[m[2]])
    if m := re.fullmatch(r"q([1-4])[\s\-']*(?:fy)?(\d{4}|\d{2})", text):
        return f"{_year(m[2]):04d}-Q{m[1]}"
    if m := re.fullmatch(r"(?:fy)?(\d{4})[\s\-]*q([1-4])", text):
        return f"{int(m[1]):04d}-Q{m[2]}"
    return None


def _year(text: str) -> int:
    year = int(text)
    return year + 2000 if year < 100 else year


def _ym(year: int, month: int) -> str | None:
    return f"{year:04d}-{month:02d}" if 1 <= month <= 12 and 1900 < year < 2200 else None


def _text(cell: Cell) -> str:
    if cell is None:
        return ""
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


# Reading files -------------------------------------------------------------


def read_grid(path: str | Path, sheet: str | None = None) -> tuple[list[list[Cell]], str | None]:
    """Load a CSV or Excel sheet as a list of rows. Picks the busiest sheet if none is named."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".csv", ".txt", ".tsv"):
        raw = path.read_text(encoding="utf-8-sig", errors="replace")
        try:
            dialect: Any = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        return [list(r) for r in csv.reader(raw.splitlines(), dialect)], None
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            names = [sheet] if sheet else wb.sheetnames
            if sheet and sheet not in wb.sheetnames:
                raise ValueError(f"No sheet named {sheet!r}. Sheets: {wb.sheetnames}")
            best: tuple[int, str, list[list[Cell]]] | None = None
            for name in names:
                rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
                numeric = sum(1 for r in rows for c in r if parse_amount(c) is not None)
                if best is None or numeric > best[0]:
                    best = (numeric, name, rows)
            assert best is not None
            return best[2], best[1]
        finally:
            wb.close()
    if suffix == ".xls":
        raise ValueError("Old .xls files aren't supported. Save the file as .xlsx or .csv first.")
    raise ValueError(f"Unsupported file type {suffix!r}. Use .csv or .xlsx.")


# Inspection ------------------------------------------------------------------

LONG_COLUMNS = {
    "account": re.compile(r"^(account( name)?|gl account|description|line( item)?|name|category)$"),
    "number": re.compile(r"^(account (no|number|#|code)|acct( no| #)?|gl( code)?|number|code|no)$"),
    "period": re.compile(r"^(period|month|date|posting period|fiscal period|quarter)$"),
    "amount": re.compile(r"^(amount|value|net|balance|net amount|actual|budget|total)$"),
    "debit": re.compile(r"^(debit|debits|dr)$"),
    "credit": re.compile(r"^(credit|credits|cr)$"),
    "department": re.compile(
        r"^(department|dept|class|cost cent(er|re)|division|location|business unit|segment)$"
    ),
}


def inspect_file(
    path: str | Path, sheet: str | None = None, period: str | None = None
) -> Inspection:
    """Work out a file's layout and return its rows, with notes on everything skipped."""
    grid, sheet_name = read_grid(path, sheet)
    units, units_note = _detect_units(grid)
    header = _find_wide_header(grid)
    if header is not None:
        insp = _inspect_wide(str(path), sheet_name, grid, header, units)
    elif (long_header := _find_long_header(grid)) is not None:
        insp = _inspect_long(str(path), sheet_name, grid, *long_header, units)
    elif (dept_header := _find_department_header(grid)) is not None:
        if not period:
            raise ValueError(
                "This file has departments or classes across the top but no periods. "
                "Say which period it covers (for example --period 2026-09)."
            )
        insp = _inspect_wide(str(path), sheet_name, grid, dept_header, units, period=period)
    else:
        raise ValueError(
            "Couldn't find a header row with periods (like 'Jan 2026' or '2026-01') or a "
            "long-format header (account, period, amount). Check the file, or convert it to "
            "the standard format: period,line,amount,department."
        )
    if units_note:
        insp.hints.append(units_note)
    return insp


def _detect_units(grid: Sequence[Sequence[Cell]]) -> tuple[Decimal, str]:
    for r, row in enumerate(grid[:12], start=1):
        for cell in row:
            if isinstance(cell, str):
                for pattern, multiplier in UNIT_PATTERNS:
                    if pattern.search(cell):
                        return multiplier, (
                            f"Row {r} says amounts are {cell.strip()!r}; "
                            f"multiplied every amount by {multiplier:,}."
                        )
    return Decimal(1), ""


def _find_wide_header(grid: Sequence[Sequence[Cell]]) -> tuple[int, dict[int, str]] | None:
    for r, row in enumerate(grid[:40]):
        cols = {c: p for c, cell in enumerate(row) if (p := parse_period(cell))}
        if len(cols) >= 2 or (len(cols) == 1 and _has_numbers_below(grid, r, next(iter(cols)))):
            return r, cols
    return None


def _has_numbers_below(grid: Sequence[Sequence[Cell]], r: int, c: int) -> bool:
    below = [row[c] for row in grid[r + 1 : r + 30] if c < len(row)]
    return sum(parse_amount(x) is not None for x in below) >= 2


def _find_long_header(grid: Sequence[Sequence[Cell]]) -> tuple[int, dict[str, int]] | None:
    for r, row in enumerate(grid[:40]):
        found: dict[str, int] = {}
        for c, cell in enumerate(row):
            text = _text(cell).lower().strip(" :#")
            for role, pattern in LONG_COLUMNS.items():
                if role not in found and pattern.match(text):
                    found[role] = c
                    break
        has_amount = "amount" in found or ("debit" in found and "credit" in found)
        if "period" in found and has_amount and ("account" in found or "number" in found):
            return r, found
    return None


def _find_department_header(grid: Sequence[Sequence[Cell]]) -> tuple[int, dict[int, str]] | None:
    for r, row in enumerate(grid[:40]):
        cols = {
            c: _text(cell)
            for c, cell in enumerate(row)
            if c > 0
            and _text(cell)
            and parse_amount(cell) is None
            and not re.fullmatch(r"total.*", _text(cell), re.IGNORECASE)
        }
        if len(cols) >= 2 and all(_has_numbers_below(grid, r, c) for c in cols):
            return r, cols
    return None


def _split_account(label: str) -> tuple[str | None, str]:
    if m := ACCOUNT_NUMBER.match(label):
        return m[1], m[2]
    return None, label


def _inspect_wide(
    path: str,
    sheet: str | None,
    grid: Sequence[Sequence[Cell]],
    header: tuple[int, dict[int, str]],
    units: Decimal,
    period: str | None = None,
) -> Inspection:
    header_row, cols = header
    first_value_col = min(cols)
    total_cols = [
        c
        for c, cell in enumerate(grid[header_row])
        if c > first_value_col and c not in cols and re.match(r"\s*total", _text(cell), re.I)
    ]
    by_department = period is not None
    insp = Inspection(
        path,
        sheet,
        "by_department" if by_department else "wide",
        header_row + 1,
        list(dict.fromkeys(cols.values())),
        units,
    )
    if total_cols:
        insp.hints.append(
            f"Ignored the total column ({_text(grid[header_row][total_cols[0]])!r}) and used it "
            "to check that each row's columns add up."
        )
    stack: list[tuple[str, int]] = []  # (section name, index into details)
    details: list[SourceRow] = []
    for r in range(header_row + 1, len(grid)):
        row = grid[r]
        label_cells = [_text(c) for c in row[:first_value_col]]
        label = " ".join(x for x in label_cells if x)
        raw_values = {cols[c]: parse_amount(row[c]) for c in cols if c < len(row)}
        values = {k: v * units for k, v in raw_values.items() if v is not None}
        n = r + 1
        if not label and not values:
            continue
        if not values:
            if label:
                stack.append((label, len(details)))
                insp.skipped.append((n, f"section header {label!r}"))
            continue
        if not label:
            insp.skipped.append((n, "numbers with no account label"))
            insp.warnings.append(f"Row {n} has amounts but no label; it was skipped.")
            continue
        _check_row_total(insp, row, total_cols, raw_values, n, label)
        number, name = _split_account(label)
        section = " / ".join(s for s, _ in stack)
        if TOTAL_WORDS.match(name):
            source = _SourceBlock.for_total(name, stack, details)
            src_row = SourceRow(n, label, number, name, values, "total", section)
            if source is not None:
                block_sum = _sum_values(source)
                if _matches(block_sum, values):
                    src_row.note = f"matches the sum of the {len(source)} rows above"
                else:
                    src_row.note = "does not match the rows above; kept for cross-checks only"
            insp.rows.append(src_row)
            insp.skipped.append((n, f"total row {label!r}"))
            continue
        block = details[stack[-1][1] :] if stack else details
        unlabeled = _unlabeled_subtotal(block, values)
        if unlabeled:
            insp.rows.append(
                SourceRow(
                    n,
                    label,
                    number,
                    name,
                    values,
                    "subtotal",
                    section,
                    note=f"equals the sum of the {unlabeled} rows above",
                )
            )
            insp.skipped.append((n, f"unlabeled subtotal {label!r}"))
            insp.hints.append(
                f"Row {n} ({label!r}) equals the sum of the {unlabeled} rows above it, so it was "
                "treated as a subtotal. Map it explicitly if that's wrong."
            )
            continue
        if by_department:
            for d, amount in values.items():
                details.append(
                    SourceRow(n, label, number, name, {period: amount}, "detail", section, d)  # type: ignore[dict-item]
                )
                insp.rows.append(details[-1])
            continue
        detail = SourceRow(n, label, number, name, values, "detail", section)
        details.append(detail)
        insp.rows.append(detail)
    if by_department:
        insp.columns = [period]  # type: ignore[list-item]
    _finish_hints(insp)
    return insp


class _SourceBlock:
    @staticmethod
    def for_total(
        name: str, stack: list[tuple[str, int]], details: list[SourceRow]
    ) -> list[SourceRow] | None:
        """Rows a 'Total X' row should equal, popping the matching section off the stack."""
        m = re.match(r"total\s+(.+)", name, re.IGNORECASE)
        if not m:
            return None  # gross profit, net income: cross-section totals, checked after mapping
        target = m[1].strip().lower()
        for i in range(len(stack) - 1, -1, -1):
            if stack[i][0].lower().strip(" :") == target:
                start = stack[i][1]
                del stack[i:]
                return details[start:]
        if stack:
            start = stack[-1][1]
            stack.pop()
            return details[start:]
        return details


def _sum_values(rows: Sequence[SourceRow]) -> dict[str, Decimal]:
    out: dict[str, Decimal] = defaultdict(lambda: ZERO)
    for row in rows:
        for k, v in row.values.items():
            out[k] += v
    return dict(out)


def _matches(a: dict[str, Decimal], b: dict[str, Decimal]) -> bool:
    keys = set(a) | set(b)
    return bool(keys) and all(abs(a.get(k, ZERO) - b.get(k, ZERO)) <= Decimal("0.5") for k in keys)


def _unlabeled_subtotal(block: Sequence[SourceRow], values: dict[str, Decimal]) -> int:
    """If the row equals the sum of every detail row in its section (2+), return how many."""
    if len(block) < 2 or not any(values.values()):
        return 0
    return len(block) if _matches(_sum_values(block), values) else 0


def _check_row_total(
    insp: Inspection,
    row: Sequence[Cell],
    total_cols: list[int],
    values: dict[str, Decimal | None],
    n: int,
    label: str,
) -> None:
    if not total_cols or total_cols[0] >= len(row):
        return
    total = parse_amount(row[total_cols[0]])
    if total is None:
        return
    row_sum = sum((v for v in values.values() if v is not None), ZERO)
    if abs(row_sum - total) > Decimal("0.5"):
        insp.warnings.append(
            f"Row {n} ({label!r}): the columns add up to {row_sum:,} but the file's total "
            f"column says {total:,}."
        )


def _inspect_long(
    path: str,
    sheet: str | None,
    grid: Sequence[Sequence[Cell]],
    header_row: int,
    cols: dict[str, int],
    units: Decimal,
) -> Inspection:
    insp = Inspection(path, sheet, "long", header_row + 1, [], units)
    grouped: dict[tuple[str, str | None, str | None], dict[str, Decimal]] = defaultdict(dict)
    first_row: dict[tuple[str, str | None, str | None], int] = {}

    def cell(row: Sequence[Cell], role: str) -> Cell:
        c = cols.get(role)
        return row[c] if c is not None and c < len(row) else None

    for r in range(header_row + 1, len(grid)):
        row = grid[r]
        n = r + 1
        if not any(_text(c) for c in row):
            continue
        period = parse_period(cell(row, "period"))
        name = _text(cell(row, "account"))
        number = _text(cell(row, "number")) or None
        if not name and number is None:
            insp.skipped.append((n, "no account"))
            continue
        if period is None:
            insp.skipped.append((n, f"unrecognized period {_text(cell(row, 'period'))!r}"))
            continue
        if "amount" in cols:
            amount = parse_amount(cell(row, "amount"))
        else:
            debit = parse_amount(cell(row, "debit")) or ZERO
            credit = parse_amount(cell(row, "credit")) or ZERO
            amount = debit - credit
        if amount is None:
            insp.skipped.append((n, "no amount"))
            continue
        if not number:
            number, name = _split_account(name)
        if TOTAL_WORDS.match(name):
            insp.skipped.append((n, f"total row {name!r}"))
            continue
        dept = _text(cell(row, "department")) or None
        label = f"{number} {name}".strip() if number else name
        key = (label, number, dept)
        grouped[key][period] = grouped[key].get(period, ZERO) + amount * units
        first_row.setdefault(key, n)
    periods: set[str] = set()
    for (label, number, dept), values in grouped.items():
        _, name = _split_account(label)
        insp.rows.append(
            SourceRow(
                first_row[(label, number, dept)],
                label,
                number,
                name,
                values,
                "detail",
                department=dept,
            )
        )
        periods.update(values)
    insp.columns = sorted(periods)
    if "debit" in cols:
        insp.hints.append(
            "Amounts are debits minus credits, so revenue will come in negative "
            "and is flipped when mapped."
        )
    _finish_hints(insp)
    return insp


def _finish_hints(insp: Inspection) -> None:
    labelled = [r for r in insp.totals if r.kind == "total"]
    mismatched = [r for r in labelled if r.note.startswith("does not match")]
    if labelled:
        insp.hints.append(
            f"Skipped {len(labelled)} total rows so nothing is counted twice; "
            f"{len(labelled) - len(mismatched)} were checked against the rows above them."
        )
    for r in mismatched:
        insp.warnings.append(
            f"Row {r.row} ({r.label!r}) doesn't equal the rows above it. Check for hidden rows "
            "or accounts outside the section."
        )
    if not insp.details:
        insp.warnings.append("No account rows were found below the header.")
