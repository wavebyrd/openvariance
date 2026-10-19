from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from openvariance import compare, demo, excel, explain
from openvariance.forecast import default_scenarios, infer_assumptions, run_scenarios
from openvariance.sec import load_company

formulas = pytest.importorskip("formulas")


def recalc(path: Path) -> dict[str, Any]:
    """Recalculate a workbook with an independent formula engine (not Excel, not our code)."""
    model = formulas.ExcelModel().loads(str(path)).finish()
    solution = model.calculate()
    out = {}
    for key, value in solution.items():
        ref = key.split("]")[-1].upper()  # SHEET!A1
        cell = value.value if hasattr(value, "value") else value
        try:
            out[ref.replace("'", "")] = cell[0][0]
        except (TypeError, IndexError):
            out[ref.replace("'", "")] = cell
    return out


def test_forecast_workbook_recalculates_to_the_engine(tmp_path: Path, demo_data: demo.Demo) -> None:
    base = infer_assumptions(demo_data.history).assumptions
    results = run_scenarios(demo_data.history, default_scenarios(base, 12), 6)
    path = excel.export_forecast(results, tmp_path / "f.xlsx", company="Demo Co.")
    values = recalc(path)
    ws = load_workbook(path)["Model"]
    rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    py = results[0]
    for i, period in enumerate(py.periods):
        col = chr(ord("C") + i)
        for label, key in (
            ("Revenue", "revenue"),
            ("Net income", "net_income"),
            ("Cash and equivalents", "cash"),
            ("Total assets", "total_assets"),
        ):
            got = Decimal(str(values[f"MODEL!{col}{rows[label]}"]))
            assert abs(got - py.values[period][key]) < Decimal("0.05"), (period, key)
        assert values[f"MODEL!{col}{rows['Cash ties to cash flow']}"] == 0
        assert values[f"MODEL!{col}{rows['Balance sheet balances (net of opening gap)']}"] == 0


def test_variance_workbook_formulas_match_the_report(tmp_path: Path, demo_data: demo.Demo) -> None:
    rep = compare(demo_data.actual, demo_data.budget, ["2026-09"])
    bridge = explain(demo_data.budget_drivers["2026-09"], demo_data.actual_drivers["2026-09"])
    path = excel.export_variance(rep, tmp_path / "v.xlsx", bridge=bridge, company="Demo Co.")
    values = recalc(path)
    ws = load_workbook(path)["Variance"]
    rows = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}
    for line in ("Operating income", "Gross profit", "Payroll and benefits"):
        r = rows[line]
        row = next(x for x in rep.rows if x.label == line)
        assert abs(Decimal(str(values[f"VARIANCE!D{r}"])) - row.impact) < Decimal("0.05")
        assert values[f"VARIANCE!F{r}"] == row.status.value
    checks = [v for k, v in values.items() if k.startswith("BRIDGE!B") and v == 0]
    assert checks, "bridge check cell should evaluate to 0"
    assert {"Dashboard", "Variance", "Bridge", "Notes"} <= set(load_workbook(path).sheetnames)


def test_revolver_formulas_match_the_engine(tmp_path: Path, crwv_facts: Path) -> None:
    history = load_company("CRWV", facts_path=crwv_facts).data
    base = infer_assumptions(history).assumptions
    results = run_scenarios(history, default_scenarios(base, 4), 4)
    py = results[0]
    assert py.values[py.periods[-1]]["revolver"] > 0  # this company needs outside funding
    path = excel.export_forecast(results, tmp_path / "crwv.xlsx")
    values = recalc(path)
    ws = load_workbook(path)["Model"]
    rows = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    for i, period in enumerate(py.periods):
        col = chr(ord("C") + i)
        for label, key in (
            ("Revolving credit facility", "revolver"),
            ("Cash and equivalents", "cash"),
            ("Net income", "net_income"),
        ):
            got = Decimal(str(values[f"MODEL!{col}{rows[label]}"]))
            assert abs(got - py.values[period][key]) < Decimal("1"), (period, key)
